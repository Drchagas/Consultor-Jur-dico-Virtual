from __future__ import annotations

import re
from datetime import date
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent.parent


def _value(row, key: str, default: str = "") -> str:
    try:
        return str(row[key] or default)
    except Exception:
        return default


def client_qualification(client) -> str:
    parts = []
    person_type = _value(client, "person_type", "Pessoa Física")
    name = _value(client, "name")
    if person_type == "Pessoa Jurídica":
        parts.append(name)
        if _value(client, "document"):
            parts.append(f"inscrita no CNPJ sob nº {_value(client,'document')}")
    else:
        parts.append(name)
        if _value(client, "nationality"):
            parts.append(_value(client, "nationality"))
        if _value(client, "marital_status"):
            parts.append(_value(client, "marital_status"))
        if _value(client, "profession"):
            parts.append(_value(client, "profession"))
        if _value(client, "rg"):
            parts.append(f"RG nº {_value(client,'rg')}")
        if _value(client, "document"):
            parts.append(f"CPF nº {_value(client,'document')}")
    address_bits = [_value(client, "address"), _value(client, "address_number"), _value(client, "complement"), _value(client, "neighborhood")]
    address = ", ".join(x for x in address_bits if x)
    city_state = "/".join(x for x in [_value(client, "city"), _value(client, "state")] if x)
    if address:
        parts.append(f"residente e domiciliado(a) em {address}" if person_type != "Pessoa Jurídica" else f"com endereço em {address}")
    if city_state:
        parts.append(city_state)
    if _value(client, "zip_code"):
        parts.append(f"CEP {_value(client,'zip_code')}")
    return ", ".join(parts) + "."


def _setup(document: Document, org) -> None:
    sec = document.sections[0]
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    styles = document.styles
    styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(11)
    logo_path = _value(org, "logo_path")
    if logo_path.startswith("/static/"):
        path = BASE_DIR / "app" / "static" / logo_path.replace("/static/", "", 1)
        if path.exists():
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try: p.add_run().add_picture(str(path), width=Cm(5.2))
            except Exception: pass


def _footer(document: Document, org) -> None:
    p = document.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details = " · ".join(x for x in [_value(org,"brand_name") or _value(org,"name"), _value(org,"address"), _value(org,"phone"), _value(org,"email")] if x)
    run = p.add_run(details); run.font.size = Pt(8)


def _title(document: Document, text: str) -> None:
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(13)


def create_power_of_attorney(client, org, target: Path, *, special_powers: Iterable[str] = (), case_number: str = "") -> None:
    doc = Document(); _setup(doc, org); _title(doc, "PROCURAÇÃO")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("OUTORGANTE: ").bold = True; p.add_run(client_qualification(client))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lawyer = _value(org, "lawyer_name") or "[CONFIGURAR ADVOGADO RESPONSÁVEL]"
    oab = _value(org, "oab_number") or "[CONFIGURAR OAB]"
    p.add_run("OUTORGADO(S): ").bold = True
    p.add_run(f"{lawyer}, advogado(a), {oab}, integrante de {_value(org,'brand_name') or _value(org,'name')}, com endereço profissional em {_value(org,'address') or '[CONFIGURAR ENDEREÇO PROFISSIONAL]'}, e-mail {_value(org,'email') or '[CONFIGURAR E-MAIL]' }.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("PODERES: ").bold = True
    base = "poderes gerais para o foro em geral, para representar o(a) outorgante judicial e extrajudicialmente, praticando os atos necessários à defesa de seus interesses"
    if case_number:
        base += f", inclusive no processo nº {case_number}"
    powers = [re.sub(r"\s+", " ", x).strip() for x in special_powers if str(x).strip()]
    if powers:
        base += ", com poderes especiais expressamente conferidos para " + ", ".join(powers)
    p.add_run(base + ".")
    doc.add_paragraph()
    city = _value(client,"city") or _value(org,"city") or "[CIDADE]"
    p = doc.add_paragraph(f"{city}, {date.today().strftime('%d/%m/%Y')}."); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph(); p = doc.add_paragraph("____________________________________________"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(_value(client,"name")); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _footer(doc, org); target.parent.mkdir(parents=True, exist_ok=True); doc.save(str(target))


def create_ajg_declaration(client, org, target: Path) -> None:
    doc = Document(); _setup(doc, org); _title(doc, "DECLARAÇÃO DE HIPOSSUFICIÊNCIA ECONÔMICA")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(client_qualification(client) + " ")
    p.add_run("DECLARA, para os fins jurídicos pertinentes, que não possui condições de arcar com custas, despesas processuais e honorários sem prejuízo de seu próprio sustento e/ou de sua família, requerendo, quando juridicamente cabível e mediante análise do advogado responsável, a concessão dos benefícios da gratuidade da justiça.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Declara, ainda, estar ciente de que as informações prestadas são de sua responsabilidade e poderão ser submetidas à comprovação quando exigida.")
    doc.add_paragraph()
    city = _value(client,"city") or "[CIDADE]"
    p = doc.add_paragraph(f"{city}, {date.today().strftime('%d/%m/%Y')}."); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph(); p = doc.add_paragraph("____________________________________________"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(_value(client,"name")); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _footer(doc, org); target.parent.mkdir(parents=True, exist_ok=True); doc.save(str(target))


# --------------------------------------------------------------------------
# JARBAS 8.9 — documentos que o escritório fazia fora do sistema.
# O JARBAS já tinha todos os dados; só não produzia o papel.
# --------------------------------------------------------------------------

def _paragrafo(document: "Document", texto: str, negrito: bool = False,
               alinhamento=None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = document.add_paragraph()
    r = p.add_run(texto)
    r.bold = negrito
    p.alignment = alinhamento if alinhamento is not None else WD_ALIGN_PARAGRAPH.JUSTIFY


def _assinaturas(document: "Document", nomes: list[str]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document.add_paragraph()
    for nome in nomes:
        p = document.add_paragraph("_" * 45)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q = document.add_paragraph(nome)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _reais(valor) -> str:
    try:
        v = float(valor or 0)
    except (TypeError, ValueError):
        v = 0.0
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def create_fee_contract(client, org, target: Path, *, contract=None,
                        case_number: str = "") -> None:
    """Contrato de honorários. A tabela fee_contracts já guardava valor,
    entrada, parcelas e êxito — mas o documento nunca era gerado."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    _setup(d, org)
    _title(d, "CONTRATO DE PRESTAÇÃO DE SERVIÇOS ADVOCATÍCIOS")

    _paragrafo(d, "CONTRATANTE: " + client_qualification(client))
    _paragrafo(d, "")
    _paragrafo(d, f"CONTRATADO: {_value(org, 'name')}, escritório de advocacia "
                  f"com endereço em {_value(org, 'address')}, neste ato representado "
                  f"pelo advogado signatário, inscrito na OAB.")
    _paragrafo(d, "")
    _paragrafo(d, "As partes têm entre si justo e contratado o seguinte:")

    valor = _value(contract, "contract_value", "0") if contract else "0"
    entrada = _value(contract, "entry_amount", "0") if contract else "0"
    parcelas = _value(contract, "installment_count", "0") if contract else "0"
    exito = _value(contract, "success_percent", "0") if contract else "0"
    objeto = _value(contract, "title") if contract else ""

    _paragrafo(d, "CLÁUSULA 1ª — DO OBJETO", negrito=True)
    _paragrafo(d, f"O CONTRATADO prestará serviços advocatícios referentes a "
                  f"{objeto or '[DESCREVER O OBJETO]'}"
                  + (f", processo nº {case_number}" if case_number else "") + ".")

    _paragrafo(d, "CLÁUSULA 2ª — DOS HONORÁRIOS", negrito=True)
    _paragrafo(d, f"Os honorários contratuais são de {_reais(valor)}"
                  + (f", com entrada de {_reais(entrada)}" if float(entrada or 0) > 0 else "")
                  + (f" e o saldo em {parcelas} parcela(s)" if int(float(parcelas or 0)) > 0 else "")
                  + ".")
    if float(exito or 0) > 0:
        _paragrafo(d, f"Além dos honorários contratuais, será devido o percentual de "
                      f"{exito}% sobre o proveito econômico obtido, a título de "
                      f"honorários de êxito.")
    _paragrafo(d, "Os honorários de sucumbência pertencem ao CONTRATADO, nos termos "
                  "do art. 23 da Lei 8.906/94, e não se confundem com os contratuais.")

    _paragrafo(d, "CLÁUSULA 3ª — DAS DESPESAS", negrito=True)
    _paragrafo(d, "Custas, taxas, emolumentos, honorários periciais e demais despesas "
                  "processuais correm por conta do CONTRATANTE e não estão incluídos "
                  "nos honorários acima.")

    _paragrafo(d, "CLÁUSULA 4ª — DAS OBRIGAÇÕES DO CONTRATANTE", negrito=True)
    _paragrafo(d, "O CONTRATANTE obriga-se a fornecer documentos e informações "
                  "verdadeiros e completos, a manter endereço e telefone atualizados "
                  "e a comparecer aos atos processuais quando intimado.")

    _paragrafo(d, "CLÁUSULA 5ª — DA REVOGAÇÃO E DA RENÚNCIA", negrito=True)
    _paragrafo(d, "Em caso de revogação do mandato ou desistência pelo CONTRATANTE, "
                  "serão devidos os honorários proporcionais aos serviços efetivamente "
                  "prestados até a data, apurados conforme a fase processual alcançada, "
                  "além do reembolso das despesas já suportadas.")

    _paragrafo(d, "CLÁUSULA 6ª — DO TRATAMENTO DE DADOS", negrito=True)
    _paragrafo(d, "O CONTRATANTE autoriza o tratamento de seus dados pessoais para a "
                  "finalidade exclusiva da prestação dos serviços aqui contratados, "
                  "nos termos da Lei 13.709/2018, inclusive mediante uso de ferramentas "
                  "de inteligência artificial e eventual transferência internacional a "
                  "provedores de tecnologia, sempre sob sigilo profissional.")

    _paragrafo(d, "CLÁUSULA 7ª — DO FORO", negrito=True)
    _paragrafo(d, f"Fica eleito o foro da comarca de {_value(org, 'city') or '[COMARCA]'} "
                  "para dirimir controvérsias, sem prejuízo do foro do domicílio do "
                  "CONTRATANTE quando aplicável a legislação consumerista.")

    _paragrafo(d, "")
    p = d.add_paragraph(f"{_value(org, 'city') or '[CIDADE]'}, "
                        f"{datetime.now().strftime('%d de %B de %Y')}.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    _assinaturas(d, [_value(client, "name") + " — CONTRATANTE",
                     _value(org, "name") + " — CONTRATADO",
                     "Testemunha 1 — CPF:",
                     "Testemunha 2 — CPF:"])
    _paragrafo(d, "As duas testemunhas acima conferem ao contrato natureza de título "
                  "executivo extrajudicial (art. 784, III, do CPC).")
    _footer(d, org)
    d.save(str(target))


def create_substabelecimento(client, org, target: Path, *,
                             substabelecido: str = "", com_reservas: bool = True,
                             case_number: str = "") -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document()
    _setup(d, org)
    tipo = "COM RESERVAS" if com_reservas else "SEM RESERVAS"
    _title(d, f"SUBSTABELECIMENTO {tipo} DE PODERES")
    _paragrafo(d, f"O advogado signatário, mandatário de "
                  f"{_value(client, 'name')}"
                  + (f", nos autos do processo nº {case_number}" if case_number else "")
                  + f", SUBSTABELECE {tipo.lower()} os poderes que lhe foram conferidos "
                  f"a {substabelecido or '[NOME E OAB DO SUBSTABELECIDO]'}.")
    if com_reservas:
        _paragrafo(d, "O substabelecimento é feito com reserva de iguais poderes, "
                      "permanecendo o subscritor no patrocínio da causa.")
    _paragrafo(d, "")
    p = d.add_paragraph(f"{_value(org, 'city') or '[CIDADE]'}, "
                        f"{datetime.now().strftime('%d de %B de %Y')}.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _assinaturas(d, ["Advogado substabelecente"])
    _footer(d, org)
    d.save(str(target))


def create_honorarios_receipt(client, org, target: Path, *, valor=0,
                              referente: str = "", forma: str = "") -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document()
    _setup(d, org)
    _title(d, "RECIBO DE HONORÁRIOS ADVOCATÍCIOS")
    _paragrafo(d, f"Recebi de {client_qualification(client)} a importância de "
                  f"{_reais(valor)}"
                  + (f", paga por {forma}" if forma else "")
                  + f", referente a {referente or '[DESCREVER]'}, pelo que firmo "
                  "o presente recibo para que produza seus efeitos legais.")
    _paragrafo(d, "")
    p = d.add_paragraph(f"{_value(org, 'city') or '[CIDADE]'}, "
                        f"{datetime.now().strftime('%d de %B de %Y')}.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _assinaturas(d, [_value(org, "name")])
    _footer(d, org)
    d.save(str(target))


def create_hipossuficiencia(client, org, target: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document()
    _setup(d, org)
    _title(d, "DECLARAÇÃO DE HIPOSSUFICIÊNCIA ECONÔMICA")
    _paragrafo(d, f"{client_qualification(client)} DECLARA, sob as penas da lei, "
                  "não possuir condições de arcar com as custas processuais e os "
                  "honorários advocatícios sem prejuízo do próprio sustento e do de "
                  "sua família, requerendo os benefícios da gratuidade da justiça, "
                  "nos termos do art. 98 do Código de Processo Civil.")
    _paragrafo(d, "Declara ainda estar ciente de que a falsidade desta declaração "
                  "sujeita o declarante às sanções dos arts. 100, parágrafo único, "
                  "do CPC e 299 do Código Penal.")
    _paragrafo(d, "")
    p = d.add_paragraph(f"{_value(org, 'city') or '[CIDADE]'}, "
                        f"{datetime.now().strftime('%d de %B de %Y')}.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _assinaturas(d, [_value(client, "name")])
    _footer(d, org)
    d.save(str(target))


MODELOS = {
    "procuracao": ("Procuração ad judicia", create_power_of_attorney),
    "ajg": ("Declaração de AJG", create_ajg_declaration),
    "contrato_honorarios": ("Contrato de honorários", create_fee_contract),
    "substabelecimento": ("Substabelecimento", create_substabelecimento),
    "recibo_honorarios": ("Recibo de honorários", create_honorarios_receipt),
    "hipossuficiencia": ("Declaração de hipossuficiência", create_hipossuficiencia),
}
